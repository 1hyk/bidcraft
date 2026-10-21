"""核对最终PDF印刷页码和Word章节引用；不判断技术内容或真实证明充分性。"""
import hashlib
import re
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def pagination_from_footers(footers, template):
    compact=lambda value:re.sub(r'\s+','',value)
    pre, token, post=compact(template).partition('{page}')
    if not token:raise ValueError('页脚缺少页码字段，无法核对最终分页')
    pattern=re.compile('^'+re.escape(pre)+r'(\d+)'+re.escape(post)+'$')
    pages=[]
    for physical,text in enumerate(footers,1):
        found={int(m[1]) for line in text.splitlines() if (m:=pattern.fullmatch(compact(line)))}
        if len(found)>1:raise ValueError('PDF页脚页码存在歧义，未采用旧分页统计')
        if found:pages.append((physical,found.pop()))
    if not pages or [n for _,n in pages]!=list(range(1,len(pages)+1)):
        raise ValueError('PDF正文页码缺失或不连续，未采用旧分页统计')
    if [p for p,_ in pages]!=list(range(pages[0][0],pages[-1][0]+1)):
        raise ValueError('PDF正文物理页范围不连续')
    return {'body_start':pages[0][0], 'body_end':pages[-1][0], 'body_pages':len(pages)},dict(pages)


def page_fields(document):
    result=[]
    for paragraph in document.element.body.xpath('.//w:p'):
        depth=0;instruction='';value='';separate=False
        for element in paragraph.iter():
            if element.tag==qn('w:fldChar'):
                kind=element.get(qn('w:fldCharType'))
                if kind=='begin':
                    depth+=1
                    if depth==1:instruction='';value='';separate=False
                elif kind=='separate' and depth==1:separate=True
                elif kind=='end':
                    if depth==1:
                        match=re.search(r'\bPAGEREF\s+bc_node_(\d+)\b',instruction)
                        if match:result.append((int(match[1]),value.strip()))
                    depth=max(0,depth-1)
            elif depth==1 and element.tag==qn('w:instrText'):instruction+=element.text or ''
            elif depth and separate and element.tag==qn('w:t'):value+=element.text or ''
    return result


def verify_references(expected, actual, headings, texts, printed):
    normalize=lambda value:re.sub(r'\s+','',value)
    result=[]
    for identity,value in actual:
        target=headings.get(identity,'')
        hits=[physical for physical,number in printed.items()
              if str(number)==value and target and normalize(target) in normalize(texts[physical-1])]
        result.append({'node_id':identity,'page':value,'physical_pages':hits,'verified':bool(hits)})
    same=[n for n,_ in expected]==[n for n,_ in actual]
    return {'valid':bool(expected) and same and all(r['verified'] for r in result),
            'references':result,'expected_count':len(expected),'actual_count':len(actual)}


def validate_files(source, output_dir, footer_format):
    reader=PdfReader(output_dir/'document.pdf')
    texts=[];footers=[]
    for page in reader.pages:
        footer=[]
        def visit(text,cm,tm,font,size):
            if 0<=tm[5]<float(page.mediabox.height)*.10:footer.append(text)
        texts.append(page.extract_text(visitor_text=visit) or '')
        footers.append(''.join(footer))
    pagination,printed=pagination_from_footers(footers,footer_format)
    original=Document(source);rendered=Document(output_dir/'document.docx')
    # 导出自带的待补图附录另起一页，不计入目标正文篇幅。
    appendix=[p.text for p in original.paragraphs if p.text.startswith('附：本稿待补图清单（共 ')]
    if appendix:
        compact=lambda value:re.sub(r'\s+','',value)
        hits=[i for i in printed if any(compact(title) in compact(texts[i-1]) for title in appendix)]
        if len(hits)!=1 or hits[0]<=pagination['body_start']:
            raise ValueError('无法确认最终PDF的正文与附录边界，未采用旧分页统计')
        pagination.update(body_end=hits[0]-1,body_pages=hits[0]-pagination['body_start'])
    headings={}
    for bookmark in rendered.element.body.xpath('.//w:bookmarkStart'):
        name=bookmark.get(qn('w:name'),'')
        if name.startswith('bc_node_') and name[8:].isdigit():
            headings[int(name[8:])]=''.join(bookmark.getparent().xpath('.//w:t/text()'))
    result=verify_references(page_fields(original),page_fields(rendered),headings,texts,printed)
    result.update(docx_sha256=digest(output_dir/'document.docx'),pdf_sha256=digest(output_dir/'document.pdf'),
                  note='仅核对当前排版文件的章节页码引用，不代表技术措施或真实材料已验收')
    return {**pagination,'total_pages':len(reader.pages),'pagination_basis':'final_pdf_footer','export_validation':result}


def current_validation(project_id, signature):
    from app.services.pagination_service import stored
    from app.services.export_dependencies import file_stamp
    from app.core.config import settings
    measurement=stored(project_id) or {}
    validation=measurement.get('export_validation') or {}
    if measurement.get('signature')!=signature or not validation.get('valid'):
        return None
    directory=settings.uploads_dir/'layout'/f'project_{int(project_id)}'/signature
    if any(not validation.get(kind+'_sha256') or file_stamp(directory/('document.'+kind))!=validation[kind+'_sha256']
           for kind in ('docx','pdf')):
        return None
    return validation
