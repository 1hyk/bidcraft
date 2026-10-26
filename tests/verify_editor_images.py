"""单图接口与导出契约，不调用模型、不写业务数据库。"""
import sys,io,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'backend'))
from fastapi import FastAPI
from fastapi.testclient import TestClient
from PIL import Image
from docx import Document
from app.api import image_assets as api
from app.core.config import settings
from app.core.docx_style import DocxStyle
from app.services.docx_export import _render_html
from app.services.image_ownership import ProjectImages

class EditorImages(unittest.TestCase):
    def setUp(self):
        app=FastAPI();app.include_router(api.router);self.client=TestClient(app)
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.directory=Path(self.temp.name)
        self.path=self.directory/'sample.png';Image.new('RGB',(100,80),'blue').save(self.path)
        self.asset={'asset_id':'IMG-'+'a'*24,'caption':'场景示意','url':'/uploads/sample.png','path':self.path,'prompt':'不要公开提示词','generated':True}
        for obj,name,value in [(api,'query_one',{'id':20}),(api,'catalog',[self.asset])]:
            mock=patch.object(obj,name,return_value=value);mock.start();self.addCleanup(mock.stop)
    def test_preview_returns_dimensions_without_local_paths(self):
        self.asset['caption']='场景示意（AI生成示意，待人工核对）'
        r=self.client.get('/api/projects/20/image-assets').json()['items'][0]
        self.assertEqual((r['width'],r['height']),(100,80));self.assertNotIn('path',r);self.assertNotIn('prompt',r)
        self.assertEqual(r['caption'],'场景示意（AI生成示意）')
    def test_regeneration_returns_candidate_and_requests_native_resolution(self):
        with patch.object(settings,'uploads_dir',self.directory),patch('app.services.section_service.generation_node_ids',return_value={706}),patch.object(settings,'image_generation_model','qwen-image-2.0'),patch.object(api.generated_images,'generate',return_value=self.asset) as generate:
            r=self.client.post('/api/projects/20/image-assets/regenerate',json={'asset_id':self.asset['asset_id'],'node_id':706,'feedback':'突出高度分层','quality':'high','request_id':'test-request-1'})
        self.assertEqual(r.status_code,200)
        need=generate.call_args.args[2]
        self.assertEqual(need['_generation_size'],'2048*1632')  # 100×80 按长边缩到 2048 并取 16 的倍数
        self.assertTrue(api.generated_images.eligible(need),'API生成的需求必须通过实际生图入口校验')
        self.assertEqual(need['_feedback'],'突出高度分层');self.assertNotIn('test-request-1',need['prompt'])
        self.assertEqual(generate.call_args.args[3],'','does not send the whole chapter')
    def test_foreign_node_rejected_before_model_call(self):
        with patch('app.services.section_service.generation_node_ids',return_value={706}),patch.object(api.generated_images,'generate') as generate:
            r=self.client.post('/api/projects/20/image-assets/regenerate',json={'asset_id':self.asset['asset_id'],'node_id':999,'feedback':'重新生成','request_id':'test-request-1'})
        self.assertEqual(r.status_code,404);generate.assert_not_called()
    def test_upload_validates_pixels_and_registers_local_asset(self):
        with patch.object(settings,'uploads_dir',self.directory):
            r=self.client.post('/api/projects/20/image-assets/upload',files={'file':('image.png',self.path.read_bytes(),'image/png')},data={'caption':'本公司实物图'})
            self.assertEqual(r.status_code,200);self.assertFalse(r.json()['generated'])
            r=self.client.post('/api/projects/20/image-assets/upload',files={'file':('fake.png',b'not png','image/png')})
            self.assertEqual(r.status_code,400)
    def test_editor_node_exports_single_image_without_ids_or_history(self):
        self.asset['caption']='场景示意（AI生成示意，待人工核对）'
        marker=f"[图片占位：场景示意|素材ID={self.asset['asset_id']}]"
        html=f'<p>前文<span data-material-image="true" data-history="old-id">{marker}</span>后文</p>'
        doc=Document();_render_html(doc,DocxStyle(),html,images=ProjectImages(assets={self.asset['asset_id']:self.asset}))
        self.assertEqual(len(doc.inline_shapes),1)
        text=''.join(p.text for p in doc.paragraphs)
        self.assertIn('前文',text);self.assertIn('后文',text);self.assertNotIn('IMG-',text);self.assertNotIn('old-id',text)
        self.assertNotIn('待人工核对',text);self.assertIn('场景示意',text);self.assertNotIn('AI生成示意',text)

    def test_custom_caption_survives_export_without_renaming_asset(self):
        self.asset['generated']=True
        marker=f"[图片占位：人工修改的图名|素材ID={self.asset['asset_id']}|自定义图注=1]"
        doc=Document();_render_html(doc,DocxStyle(),'<p>'+marker+'</p>',images=ProjectImages(assets={self.asset['asset_id']:self.asset}))
        text=''.join(p.text for p in doc.paragraphs)
        self.assertIn('人工修改的图名',text)
        self.assertNotIn('人工修改的图名（AI生成示意）',text)
        self.assertNotIn('AI生成示意',text)
        self.assertNotIn('自定义图注',text);self.assertNotIn('IMG-',text)
        self.assertEqual(self.asset['caption'],'场景示意')

    def test_schedule_and_data_chart_never_request_raster_generation(self):
        from app.services.image_planning import normalize_outline
        from app.services.generated_images import eligible
        for name in ['7*24小时值班排班示意图','监视数据准确性评估对比图表']:
            need={'name':name,'source':'illustration','purpose':'说明方案'}
            self.assertFalse(eligible(need))
            result=normalize_outline({'image_needs':[need]})
            self.assertEqual(result['image_needs'],[])
            self.assertEqual(result['table_needs'][0]['name'],name)

    def test_scene_regeneration_retains_original_context(self):
        self.asset['generation_need']={'purpose':'局部场地的跑道与机巢布置','prompt':'包括跑道、机巢和垂直高度层'}
        with patch.object(settings,'uploads_dir',self.directory),patch('app.services.section_service.generation_node_ids',return_value={706}),patch.object(settings,'image_generation_model','qwen-image-2.0'),patch.object(api.generated_images,'generate',return_value=self.asset) as generate:
            r=self.client.post('/api/projects/20/image-assets/regenerate',json={'asset_id':self.asset['asset_id'],'node_id':706,'feedback':'减少背景细节','quality':'high','request_id':'test-request-2'})
        self.assertEqual(r.status_code,200)
        self.assertIn('垂直高度层',generate.call_args.args[2]['prompt'])
        self.assertEqual(generate.call_args.args[2]['_feedback'],'减少背景细节')

if __name__=='__main__':unittest.main()
