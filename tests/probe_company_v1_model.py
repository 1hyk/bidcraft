"""可选真实模型探针：合成资料及连接级临时表，不导入真实公司库。"""
import sys,json,tempfile,re
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'backend'))
from app.core.database import get_conn,execute,query_one,query
from app.services import company_materials as m,company_fill as f


def main():
    if '--live' not in sys.argv:
        print('显式加 --live 才调用已配置模型');return
    from docx import Document
    tables=('company_profile','company_document','company_fact','section_content','company_fill_revision')
    conn=get_conn()
    try:
        with conn.cursor() as cur:
            for table in tables:
                cur.execute('SHOW CREATE TABLE '+table)
                schema=cur.fetchone()['Create Table'].replace('CREATE TABLE','CREATE TEMPORARY TABLE',1)
                cur.execute(re.sub(r',\n\s*CONSTRAINT[^\n]+','',schema))
        with tempfile.TemporaryDirectory() as tmp,patch.object(m,'storage',return_value=Path(tmp)),patch.object(f,'allowed'):
            execute('INSERT INTO company_profile (id,content) VALUES (1,%s)',(json.dumps({'name':'示例星河软件有限公司'},ensure_ascii=False),))
            path=Path(tmp)/'合成验证样本.docx';doc=Document()
            doc.add_paragraph('公司名称：示例星河软件有限公司。')
            doc.add_paragraph('资质名称：示例软件质量服务证书。证书编号：ZX-4712。持证单位：示例星河软件有限公司。有效期至2028年12月31日。')
            doc.save(path)
            execute("INSERT INTO company_document (id,name,sha256,path,status) VALUES (1,'合成验证样本.docx',%s,%s,'extracting')",('b'*64,str(path)))
            m.extract_task(0,1)
            candidates=query('SELECT * FROM company_fact')
            assert candidates and all(x['status']=='pending' for x in candidates),'提取不得自动核实'
            fact=next((x for x in candidates if x['value']=='ZX-4712'),None)
            assert fact,'模型未按原文提取随机证书编号'
            assert fact['value'] in fact['evidence']
            execute("UPDATE company_fact SET status='confirmed' WHERE id=%s",(fact['id'],))
            payload={'format':'html','body':'<p>本公司软件质量服务证书编号：[待补充：公司证书编号]</p><p>本次负责人：[待确认：项目负责人]</p>'}
            execute('INSERT INTO section_content (chapter_node_id,content,ai_generated) VALUES (999,%s,1)',(json.dumps(payload,ensure_ascii=False),))
            suggestions=f.suggest(999,999)
            print('synthetic_suggestions',json.dumps([{k:x.get(k) for k in ('placeholder','status','replacement','message')} for x in suggestions['items']],ensure_ascii=False))
            chosen=[x for x in suggestions['items'] if x['status']=='suggested']
            assert len(chosen)==1 and chosen[0]['replacement']=='ZX-4712','事实建议未准确命中'
            assert any(x['status']=='decision' for x in suggestions['items']),'本次任命未被阻断'
            result=f.apply(999,999,[chosen[0]['token']])
            body=m.decode(query_one('SELECT content FROM section_content WHERE chapter_node_id=999')['content'],{})['body']
            assert 'ZX-4712' in body and '[待确认：项目负责人]' in body
            f.undo(999,result['revision_id'])
            assert m.decode(query_one('SELECT content FROM section_content WHERE chapter_node_id=999')['content'],{})==payload
            report={'synthetic_only':True,'candidate_count':len(candidates),'random_certificate_extracted':True,'source_verified':True,'suggestion_matched':True,'appointment_blocked':True,'applied_and_undone':True}
            output=Path(__file__).parent/'output';output.mkdir(exist_ok=True)
            (output/'company_v1_model_probe.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
            print(json.dumps(report,ensure_ascii=False))
    finally:
        conn.rollback()
        with conn.cursor() as cur:
            for table in reversed(tables):cur.execute('DROP TEMPORARY TABLE IF EXISTS '+table)


if __name__=='__main__':main()
