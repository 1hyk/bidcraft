"""评分索引预览与导出一致性，不调用模型、不修改项目。"""
import sys
import unittest
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'backend'))
from app.services.scoring_index import build_entries, index_required, project_index, editable_html, criteria_summary
from app.services.docx_export import _render_editor_indexes, _walk_nodes
from app.core.docx_style import DocxStyle
from docx import Document


class ScoringIndexTests(unittest.TestCase):
    def test_only_explicit_requirement_enables_index(self):
        def row(content, category='编制要求', **kw):
            return dict(content=content, category=category, **kw)
        self.assertTrue(index_required([row('响应人须根据评分细则编制索引目录（填写对应证明材料的所在页码）。')]))
        self.assertTrue(index_required([row('投标人应提供技术响应索引表。')]))
        for rows in ([], [row('附件：评分表，序号、分值、页码')],
                     [row('无需编制索引目录')], [row('可不提供响应索引表')],
                     [row('应编制索引目录', status='ignored')],
                     [row('须提供索引表', category='评分办法')]):
            self.assertFalse(index_required(rows))

    def test_scoring_items_alone_do_not_export_index(self):
        with patch('app.services.requirement_service.list_requirements', return_value=[
            {'category':'评分办法', 'title':'维护保障'}]), patch('app.services.section_service.get_tree') as tree:
            self.assertEqual(project_index(8), [])
            tree.assert_not_called()

    def test_nested_and_unmatched_and_ignored(self):
        rows = build_entries([
            {'id': 1, 'title': '维护保障', 'content': '原文'},
            {'id': 2, 'title': '价格评分'},
            {'id': 3, 'title': '人员计划', 'status': 'ignored'},
            {'id': 4, 'title': '组织安排', 'suggested_chapter': '人员计划'},
        ], [{'id': 10, 'title': '方案', 'children': [
            {'id': 11, 'title': '维护保障方案'}, {'id': 12, 'title': '人员计划'}]}])
        self.assertEqual([r['requirement_id'] for r in rows], [1, 2, 4])
        self.assertEqual([r['node_id'] for r in rows], [11, None, 12])
        self.assertEqual(rows[0]['content'], '原文')

    def test_export_uses_preview_data_and_bookmarks(self):
        entries = [{'title': '非低空评分项', 'node_id': 912, 'chapter_title': '维护安排', 'included': True},
                   {'title': '未匹配项', 'node_id': None, 'chapter_title': '', 'included': False}]
        doc = Document()
        tree=[{'id':1,'title':'技术方案','children':[
            {'id':2,'title':'技术评分响应索引','outline':{'_document_role':'scoring_index'}},
            {'id':912,'title':'维护安排'}]}]
        content=editable_html(entries).replace('非低空评分项','人工编辑后的评分项')
        contents={2:{'format':'html','body':content},912:{'format':'html','body':'<p>维护正文</p>'}}
        _render_editor_indexes(doc, DocxStyle(), tree, contents)
        _walk_nodes(doc, DocxStyle(), tree, contents, 1)
        self.assertEqual(len(doc.tables), 1)
        self.assertEqual(len(doc.tables[0].rows), 2)
        self.assertEqual(doc.tables[0].cell(1, 0).text, '人工编辑后的评分项')
        self.assertEqual(doc.tables[0].cell(1, 2).text, entries[0]['chapter_title'])
        self.assertIn('PAGEREF bc_node_912', doc.tables[0].cell(1, 3)._tc.xml)
        self.assertFalse(doc.tables[0].autofit)
        widths=[c.width for c in doc.tables[0].columns]
        self.assertGreater(widths[1],widths[0])
        self.assertGreater(widths[1],widths[3])

    def test_criteria_and_nested_chapters_are_kept_without_inventing_response(self):
        rows=build_entries([{'id':9,'title':'维护保障','content':'应说明巡检频次；无方案得0分。'}],
            [{'id':10,'title':'维护保障','children':[{'id':11,'title':'巡检安排'},{'id':12,'title':'应急处置'}]}])
        self.assertEqual([r['id'] for r in rows[0]['chapter_refs']],[10,11,12])
        html=editable_html(rows)
        self.assertIn('应说明巡检频次；无方案得0分。',html)
        self.assertIn('维护保障',html)
        self.assertNotIn('data-page-reference="11"',html)
        self.assertNotIn('完全满足',html)

    def test_summary_keeps_requirements_and_removes_score_bands(self):
        text='方案必须包括交付批次规划、验证流程、问题处理机制。7–10分：完整合理；3–7分：存在风险。'
        self.assertEqual(criteria_summary(text),'方案必须包括交付批次规划、验证流程、问题处理机制。')
        penalty='关键项每偏离一项扣4分，其他条款每偏离一项扣2分。'
        self.assertEqual(criteria_summary(penalty),penalty)

    def test_stale_reference_is_not_exported_as_valid_field(self):
        doc = Document()
        tree=[{'id':2,'title':'响应索引','outline':{'_document_role':'scoring_index'}}]
        content={2:{'format':'html','body':'<p><span data-page-reference="999">排版后生成</span></p>'}}
        _render_editor_indexes(doc, DocxStyle(), tree, content)
        self.assertIn('章节引用失效', ''.join(p.text for p in doc.paragraphs))
        self.assertNotIn('PAGEREF bc_node_999', doc.element.xml)

    def test_no_stored_index_means_no_hidden_index(self):
        doc=Document()
        _render_editor_indexes(doc, DocxStyle(), [{'id':1,'title':'维护方案'}], {})
        self.assertEqual(len(doc.tables),0)
        self.assertEqual(len(doc.paragraphs),0)


if __name__ == '__main__':
    unittest.main()
