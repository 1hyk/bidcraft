"""明确图题只跟随相邻图片，导出到图下方居中。"""
import sys
from pathlib import Path
from unittest.mock import patch
import tempfile
from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'backend'))
from app.services import docx_export as dx
from app.core.docx_style import DocxStyle

for before in (False, True):
    doc = Document()
    figure = '<pre><code class="language-mermaid">flowchart TD\nA--&gt;B</code></pre>'
    caption = '<p><strong>图：实质性条款响应逻辑流程图</strong></p>'
    html = (caption + figure) if before else (figure + caption)
    with patch.object(dx, '_render_mermaid', side_effect=lambda d,s,c: d.add_paragraph('[图片]')):
        dx._render_html(doc, DocxStyle(), html + '<p>普通正文</p>')
    assert [p.text for p in doc.paragraphs] == ['[图片]', '实质性条款响应逻辑流程图', '普通正文']
    cap = doc.paragraphs[1]
    assert cap.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert cap.paragraph_format.first_line_indent == 0
doc = Document()
dx._render_html(doc, DocxStyle(), '<p>图：这不是紧邻图片的段落</p><p>正文</p>')
assert doc.paragraphs[0].text == '图：这不是紧邻图片的段落'
print('PASS: figure captions below image and centered; unrelated prose preserved')

with tempfile.TemporaryDirectory() as folder:
    path=Path(folder)/'history.png'
    Image.new('RGB',(80,40),'blue').save(path)
    doc=Document()
    dx._render_html(doc,DocxStyle(),'<p>前文</p><p>[图片占位：历史测试图]</p><p>后文</p>',images={'历史测试图':path})
    image_index=next(i for i,p in enumerate(doc.paragraphs) if p._p.xpath('.//w:drawing'))
    cap=doc.paragraphs[image_index+1]
    assert cap.text=='历史测试图' and cap.alignment==WD_ALIGN_PARAGRAPH.CENTER
    assert cap.paragraph_format.first_line_indent==0
    assert doc.paragraphs[0].text=='前文' and doc.paragraphs[-1].text=='后文'
print('PASS: standalone historical image has separate centered caption below')

for before in (True,False):
    doc=Document()
    figure='<pre><code class="language-mermaid">flowchart LR\nA--&gt;B</code></pre>'
    captions=['<p>图：交付批次甘特图</p>','<p>图：批次流转图</p>']
    html=''.join(c+figure if before else figure+c for c in captions)
    with patch.object(dx,'_render_mermaid',side_effect=lambda d,s,c:d.add_paragraph('[图片]')):
        dx._render_html(doc,DocxStyle(),html)
    assert [p.text for p in doc.paragraphs]==['[图片]','交付批次甘特图','[图片]','批次流转图']
    assert all(p.alignment==WD_ALIGN_PARAGRAPH.CENTER for p in doc.paragraphs if p.text!='[图片]')
print('PASS: consecutive figures retain their own before/after captions')

doc = Document()
with patch.object(dx, '_render_mermaid', side_effect=lambda d,s,c: d.add_paragraph('[图片]')):
    dx._render_html(doc, DocxStyle(), '<h3>3.1 测试流程设计</h3><pre><code class="language-mermaid">flowchart TD\nA--&gt;B</code></pre>', base_level=3, drop_title='初验阶段实施方案')
assert doc.paragraphs[-1].text == '测试流程设计流程图'
assert doc.paragraphs[-1].alignment == WD_ALIGN_PARAGRAPH.CENTER
print('PASS: missing Mermaid caption derived from local heading')

# 导出件是给客户看的正式稿：「AI生成示意」是内部生成方式标记，图注和占位里都不该出现。
from app.services.image_caption import print_caption
assert print_caption('现场布置示意图（AI生成示意）') == '现场布置示意图'
assert print_caption('现场布置示意图（AI生成示意，待人工核对）') == '现场布置示意图'
assert print_caption('数据曲线（AI生成示意，数据仅作示例，非实测结果）') == '数据曲线（数据仅作示例，非实测结果）'
assert print_caption('现场布置示意图（ AI 生成示意 ）') == '现场布置示意图'
print('PASS: print caption drops AI-generation marker variants')

for before in (True, False):
    doc = Document()
    figure = '<pre><code class="language-mermaid">flowchart TD\nA--&gt;B</code></pre>'
    caption = '<p data-figure-caption="true">大安机场现场布置示意图（AI生成示意）</p>'
    html = (caption + figure) if before else (figure + caption)
    with patch.object(dx, '_render_mermaid', side_effect=lambda d,s,c: d.add_paragraph('[图片]')):
        dx._render_html(doc, DocxStyle(), html)
    assert [p.text for p in doc.paragraphs] == ['[图片]', '大安机场现场布置示意图'], [p.text for p in doc.paragraphs]
print('PASS: figure captions exported without AI-generation marker')

doc = Document()
dx._render_html(doc, DocxStyle(), '<p>正文</p><p>[图片占位：大安机场布置示意图（AI生成示意）|素材ID=IMG-0123456789abcdef01234567]</p>',
                images={})
assert 'AI生成示意' not in doc.paragraphs[1].text, doc.paragraphs[1].text
assert '大安机场布置示意图' in doc.paragraphs[1].text
print('PASS: unresolved image placeholder exported without AI-generation marker')
