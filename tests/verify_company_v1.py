"""公司事实、证据与补全的正反用例；不调用真实模型和真实资料。"""
import sys,json
import unittest
import tempfile
from pathlib import Path
from datetime import date,timedelta
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'backend'))
from app.services import company_materials as m, company_fill as fill


class CompanyV1(unittest.TestCase):
    def test_editor_save_keeps_source_metadata(self):
        from app.api.sections import put_content,ContentBody
        payload={'format':'html','body':'原文','used':{'pages':[3]},'prev_content':{'body':'更早正文'}}
        with patch('app.api.sections.query_one',side_effect=[{'id':1,'project_id':1,'title':'测试章节'},{'content':json.dumps(payload),'ai_generated':0},{'updated_at':'now'}]),patch('app.api.sections.execute') as write:
            put_content(1,ContentBody(format='html',body='人工修改'))
            stored=json.loads(write.call_args_list[0].args[1][0])
            self.assertEqual(stored['used'],payload['used'])
            self.assertEqual(stored['prev_content'],payload['prev_content'])
            self.assertEqual(stored['body'],'人工修改')
    def test_docx_paragraph_and_table_anchors(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'sample.docx';doc=Document();doc.add_paragraph('公司名称：示例公司');table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='编号';table.cell(0,1).text='ZX-4712';doc.save(path)
            parts=m.fragments(path)
            self.assertEqual(len(parts),2)
            self.assertIn('ZX-4712',parts[1]['text'])
            self.assertNotEqual(parts[0]['location'],parts[1]['location'])

    def test_candidate_must_have_exact_evidence(self):
        parts=[{'location':'第1页','text':'示例公司资质编号ZX-4712，有效期2028年12月31日。'}]
        good={'fragment':0,'category':'资质','subject':'示例公司','field_name':'资质编号','value':'ZX-4712','evidence':parts[0]['text'],'valid_until':'2028-12-31'}
        self.assertEqual(m.valid_candidates({'facts':[good]},parts)[0]['valid_until'],date(2028,12,31))
        for changed in ({'value':'编造编号'},{'evidence':'不存在的证据'},{'fragment':9},{'fragment':True}):
            self.assertEqual(m.valid_candidates({'facts':[{**good,**changed}]},parts),[])

    def test_expired_unverified_and_conflicting_facts_are_excluded(self):
        fact={'id':1,'category':'资质','subject':'示例','field_name':'证书编号','value':'A','status':'confirmed','revision':1}
        self.assertTrue(m.usable(fact))
        for change in ({'status':'pending'},{'status':'disabled'},{'status':'rejected'},{'valid_until':date.today()-timedelta(days=1)},{'doc_id':1,'source_status':'failed'}):
            self.assertFalse(m.usable({**fact,**change}))
        certificates=m.mark_conflicts([fact,{**fact,'id':2,'value':'B'}])
        self.assertTrue(all(m.usable(f) for f in certificates))
        identity={**fact,'field_name':'统一社会信用代码'}
        rows=m.mark_conflicts([identity,{**identity,'id':2,'value':'B'}])
        self.assertTrue(all(not m.usable(f) for f in rows))

    def test_signed_token_cannot_be_tampered(self):
        with tempfile.TemporaryDirectory() as tmp,patch.object(m,'storage',return_value=Path(tmp)):
            token=fill.sign({'created':__import__('time').time(),'fact':2})
            self.assertEqual(fill.verify(token)['fact'],2)
            with self.assertRaises(ValueError):fill.verify(token+'x')

    def payload(self):return {'format':'html','body':'<p>资质编号：[待补充：证书编号]</p>','used':{'pages':[1]}}
    def item(self,payload):
        marker=fill.MARKER.search(payload['body'])
        return {'hash':fill.digest(payload),'start':marker.start(),'end':marker.end(),'placeholder':marker.group(),'fact':1,'revision':1}
    def fact(self):return {'id':1,'value':'A&B<123>','status':'confirmed','revision':1}

    def test_replacement_preserves_payload_and_escapes_html(self):
        payload=self.payload();after,_=fill.replacement_plan(payload,[self.item(payload)],{1:self.fact()})
        self.assertIn('A&amp;B&lt;123&gt;',after['body'])
        self.assertEqual(after['used'],payload['used'])
        self.assertIn('[待补充',payload['body'])

    def test_changed_body_fact_and_duplicate_selection_are_rejected(self):
        payload=self.payload();item=self.item(payload)
        for target,items,facts in [({**payload,'body':'人工修改'},[item],{1:self.fact()}),(payload,[item],{1:{**self.fact(),'revision':2}}),(payload,[item,item],{1:self.fact()})]:
            with self.assertRaises(ValueError):fill.replacement_plan(target,items,facts)

    def test_project_appointment_is_never_filled_from_company_pool(self):
        payload={'format':'md','body':'[待补充：本次项目负责人]'}
        with self.assertRaises(ValueError):fill.replacement_plan(payload,[self.item(payload)],{1:self.fact()})
        payload={'format':'html','body':'<p>本次项目负责人：[待补充：姓名]</p>'}
        with self.assertRaises(ValueError):fill.replacement_plan(payload,[self.item(payload)],{1:self.fact()})


if __name__=='__main__':unittest.main()
