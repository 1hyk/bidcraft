"""检查实际 Word 嵌图尺寸；三种入口都必须遵守当前项目边距。"""
import sys
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path
from unittest.mock import patch

from PIL import Image
from docx import Document
import lxml.html

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'backend'))
from app.core.docx_style import DocxStyle
from app.services import docx_export as dx


class ImageLayout(unittest.TestCase):
    def test_every_image_entry_respects_page_bounds_and_aspect(self):
        with tempfile.TemporaryDirectory() as tmp:
            for size in [(1000, 300), (150, 2000)]:
                path = Path(tmp) / 'image.png'
                Image.new('RGB', size, 'white').save(path)
                for style in [DocxStyle(), replace(DocxStyle(), margin_left_cm=5.08, margin_right_cm=5.08)]:
                    for route in ['mermaid', 'placeholder', 'editor']:
                        doc = Document()
                        dx._setup_page(doc.sections[0], style)
                        if route == 'mermaid':
                            with patch('app.services.mermaid_render.render_cached', return_value=path):
                                dx._render_mermaid(doc, style, 'flowchart TD\n A-->B')
                        elif route == 'placeholder':
                            with patch.object(dx, '_lookup_name', return_value=('真实图片', path)):
                                dx._add_text(doc.add_paragraph(), style, '[图片占位：真实图片]', False, False, False, True, images={})
                        else:
                            with patch.object(dx, '_local_upload_path', return_value=path):
                                dx._add_image(doc, None, lxml.html.fromstring('<img src="/uploads/test.png">'), style)
                        output = Path(tmp) / 'result.docx'
                        doc.save(output)
                        saved = Document(output)
                        self.assertEqual(len(saved.inline_shapes), 1, route)
                        shape = saved.inline_shapes[0]
                        self.assertLessEqual(shape.width.cm, min(14, 21-style.margin_left_cm-style.margin_right_cm)+0.001, route)
                        self.assertLessEqual(shape.height.cm, 19.001, route)
                        self.assertAlmostEqual(shape.height / shape.width, size[1]/size[0], places=4)

    def test_missing_image_size_uses_current_column_width(self):
        style = replace(DocxStyle(), margin_left_cm=5.08, margin_right_cm=5.08)
        self.assertAlmostEqual(dx._fit_width_cm(Path('missing.png'), style), 10.84)

    def test_mermaid_keeps_context_space_and_single_image_line(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'tall.png'
            Image.new('RGB',(200,1800),'white').save(path)
            doc=Document()
            with patch('app.services.mermaid_render.render_cached',return_value=path):
                dx._render_mermaid(doc,DocxStyle(line_spacing=2),'flowchart TD\nA-->B')
            self.assertAlmostEqual(doc.inline_shapes[0].height.cm,17,places=3)
            self.assertEqual(doc.paragraphs[0].paragraph_format.line_spacing,1.0)

    def test_flowchart_layout_changes_invalidate_cache(self):
        from app.services import mermaid_render as mr
        old=mr._cache_path('flowchart TD\nA-->B')
        with patch.object(mr,'_FLOWCHART',{**mr._FLOWCHART,'rankSpacing':50}):
            self.assertNotEqual(old,mr._cache_path('flowchart TD\nA-->B'))
        page=mr._page(['flowchart TD\nA-->B'])
        self.assertNotIn('__FLOWCHART__',page)
        self.assertIn('"rankSpacing": 20',page)


if __name__ == '__main__':
    unittest.main()
