"""只读核对独立回归稿：素材实际落图、重复像素、标题与题注。"""
import hashlib
import io
import json
import re
import sys
from collections import Counter
from pathlib import Path
from zipfile import ZipFile
from PIL import Image, ImageOps
from docx import Document

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'backend'))
from app.services.image_retrieval import catalog,selected_ids
from app.services.proposal_quality import chapters
from app.services.pagination_service import view
from app.core.database import query


def pixels(source):
    with Image.open(source) as im:
        im=ImageOps.exif_transpose(im).convert('RGBA')
        return hashlib.sha256(str(im.size).encode()+im.tobytes()).hexdigest()


if __name__=='__main__':
    pid=int(sys.argv[1]);folder=Path(sys.argv[2]);rows=list(chapters(pid))
    assets=[a for a in catalog(pid) if not a.get('generated')]
    identities={a['asset_id'] for a in assets}
    source_pixels={pixels(a['path']):a['caption'] for a in assets}
    measurement=view(pid)
    docx_path=ROOT/measurement['docx_url'].removeprefix('/')
    exported=[]
    with ZipFile(docx_path) as z:
        for name in z.namelist():
            if name.startswith('word/media/'):
                exported.append(pixels(io.BytesIO(z.read(name))))
    document=Document(docx_path)
    headings=[p.text for p in document.paragraphs if p.style.name.startswith('Heading')]
    captions=[p.text for p in document.paragraphs if p.style.name=='Caption']
    selected={i for r in rows for i in selected_ids(r['body'])}
    used=[im for r in rows for im in (r['payload'].get('used',{}).get('images') or [])]
    nodes=query('SELECT id,parent_id,title,sort_order FROM chapter_node WHERE project_id=%s ORDER BY sort_order,id',(pid,))
    parent=next((n for n in nodes if n['title']=='项目实施方案'),None)
    phases=[n['title'] for n in nodes if parent and n['parent_id']==parent['id']]
    result={'project_id':pid,'measurement':measurement,'written_nodes':len(rows),
            'historical_candidates':len(assets),'historical_selected_ids':sorted(selected & identities),
            'historical_selected_count':len(selected & identities),
            'historical_embedded_count':len(set(exported)&set(source_pixels)),
            'historical_embedded_names':[source_pixels[h] for h in set(exported)&set(source_pixels)],
            'embedded_image_count':len(exported),'unique_image_count':len(set(exported)),
            'duplicate_image_hashes':[h for h,n in Counter(exported).items() if n>1],
            'mermaid_source_count':sum(r['body'].count('```mermaid') for r in rows),
            'body_placeholders':[(r['title'],re.findall(r'\[图片占位[^\]]*\]',r['body'])) for r in rows if '[图片占位' in r['body']],
            'used_image_count':len(used),'headings':headings,'caption_style_count':len(captions),
            'implementation_phases':phases}
    (folder/'history_audit.json').write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
    (folder/'audit_bodies.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
    print(json.dumps({k:v for k,v in result.items() if k not in {'headings','historical_selected_ids','measurement'}},ensure_ascii=False))
