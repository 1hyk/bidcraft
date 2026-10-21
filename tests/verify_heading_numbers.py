"""导出目录编号：分级、同级重置、人工编号保护及正文不变。"""
import sys
from pathlib import Path
from io import BytesIO
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'backend'))
from app.services.docx_export import _walk_nodes, _numbered_title
from app.core.docx_style import DocxStyle

for title in ['第四章 技术方案', '一、技术方案', '（二）实施方案', '1.2 测试方案', '3、服务方案']:
    assert _numbered_title(title, 1, 1) == title
assert _numbered_title('“测试”方案', 1, 10) == '十、“测试”方案'
tree = [{'id': 1, 'title': '技术方案', 'children': [
    {'id': 2, 'title': '测试方案', 'children': [{'id': 3, 'title': '功能测试'}]},
    {'id': 4, 'title': '服务方案', 'children': [{'id': 5, 'title': '响应机制'}]},
]}]
contents = {3: {'body': '正文原文', 'format': 'md'}}
doc = Document()
_walk_nodes(doc, DocxStyle(), tree, contents, 1)
buffer = BytesIO()
doc.save(buffer)
buffer.seek(0)
loaded = Document(buffer)
headings = [p.text for p in loaded.paragraphs if p.style.name.startswith('Heading')]
assert headings == ['一、技术方案', '（一）测试方案', '1.功能测试', '（二）服务方案', '1.响应机制'], headings
assert '正文原文' in [p.text for p in loaded.paragraphs]
assert tree[0]['title'] == '技术方案'
print('PASS: exported heading numbers, sibling reset, existing prefixes, unchanged content')

from app.services.docx_export import _render_html
doc = Document()
_render_html(doc, DocxStyle(), '<h2>2.4 阶段性交付物清单</h2><h3>2.4.1 批次规划</h3><h4>2.4.1.1 需求理解</h4><h4>2.4.1.2 响应方案</h4><h3>2.4.2 验收</h3>', base_level=3, drop_title='阶段性交付物清单')
assert [(p.style.name,p.text) for p in doc.paragraphs] == [
    ('Heading 4','（1）批次规划'), ('Heading 5','①需求理解'),
    ('Heading 5','②响应方案'), ('Heading 4','（2）验收')]
print('PASS: duplicate numbered title removed; nested levels retained; legacy numbers replaced')

# 父章内的小标题不能与目录中的正式子章争用编号；各层级都适用。
tree=[{'id':1,'title':'技术方案','children':[{'id':2,'title':'设备方案','children':[{'id':3,'title':'检测设备'}]}, {'id':4,'title':'服务方案'}]}]
contents={1:{'format':'html','body':'<h2>一、总体思路</h2><p>父章原文保留</p>'},
          2:{'format':'html','body':'<h3>1.响应索引与落实路径</h3><p>索引原文保留</p>'},
          3:{'format':'html','body':'<h2>### 1. 工作步骤</h2><h4>（1）### 2. 核验</h4><h2>### 2. 记录</h2><p>叶子原文保留</p>'}}
doc=Document();_walk_nodes(doc,DocxStyle(),tree,contents,1)
headings=[(p.style.name,p.text) for p in doc.paragraphs if p.style.name.startswith('Heading')]
assert headings==[('Heading 1','一、技术方案'),('Heading 2','（一）设备方案'),
                  ('Heading 3','1.检测设备'),('Heading 4','（1）工作步骤'),
                  ('Heading 5','①核验'),('Heading 4','（2）记录'),('Heading 2','（二）服务方案')],headings
for text in ['总体思路','响应索引与落实路径','父章原文保留','索引原文保留','叶子原文保留']:
    assert text in [p.text for p in doc.paragraphs]
assert contents[3]['body'].startswith('<h2>### 1.'),'不得改数据库正文对象'
print('PASS: auxiliary parent headings, skipped levels, Markdown leftovers and original content preservation')
