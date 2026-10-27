"""编辑预览和 Word 导出共用的评分索引；不代表响应充分性结论。"""
from app.core.chapter_title import related, normalize
import re
import json
from html import escape

ROLE = 'scoring_index'


def is_index(node):
    outline = node.get('outline') or {}
    if isinstance(outline, str):
        try: outline = json.loads(outline)
        except (TypeError, ValueError): return False
    return outline.get('_document_role') == ROLE


def criteria_summary(content):
    """只剔除要求之后的评分档位，不机械截断技术要求。"""
    text = str(content or '').strip()
    band = re.search(r'\d+\s*[–—~-]\s*\d+\s*分[：:（(]', text)
    if band and band.start() > 10:
        return text[:band.start()].strip()
    band = re.search(r'[。\n](?=[^。\n]{0,80}得\s*\d+\s*[-–—~]?\s*\d*\s*分)', text)
    return text[:band.start()+1].strip() if band and band.start() > 10 else text


def editable_html(entries):
    rows = []
    for entry in entries:
        if not entry['included']: continue
        # 一个评分项先定位完整响应章；不再把整棵子树铺满索引。
        refs = [{'id':entry['node_id'], 'title':entry['chapter_title']}]
        titles = ''.join('<p>'+escape(ref['title'])+'</p>' for ref in refs)
        pages = ''.join(f'<p><span data-page-reference="{int(ref["id"])}">排版后生成</span></p>' for ref in refs)
        criteria = escape(criteria_summary(entry.get('content'))).replace('\n','<br>')
        rows.append('<tr><td><p>'+escape(entry['title'])+'</p></td><td><p>'+criteria+
                    '</p></td><td>'+titles+'</td><td>'+pages+'</td></tr>')
    return ('<p>本表按技术评分要点定位响应章节；评分标准以招标原文为准，实际证明材料须据实提供。</p>'
            '<table><tr><th colwidth="130"><p>评分项</p></th><th colwidth="340"><p>评分要求摘要</p></th><th colwidth="180"><p>对应章节</p></th><th colwidth="100"><p>正文页码</p></th></tr>'+
            ''.join(rows)+'</table>')


def ensure_editor_index(project_id):
    """评分索引节点不再自动创建（2026-10-27 确认不需要）。

    接口保留为空操作，只为兼容尚未刷新的旧前端；不再读写目录树。
    """
    return {'created': False}


def _text_requires_index(text):
    """只认明确要求编制/填写索引的句子；评分表、附件名称本身不构成指令。"""
    for sentence in re.split(r'[。；;\n]', str(text or '')):
        if re.search(r'无需|无须|不需|不用|不要求|可不|自行决定|自愿', sentence):
            continue
        if re.search(
                r'(?:须|应当|应|必须|要求|请)[^。；;\n]{0,50}'
                r'(?:编制|提供|附|填写|建立|设置|列明)'
                r'[^。；;\n]{0,25}(?:索引目录|评分索引|响应索引|索引表|索引(?!号))',
                sentence):
            return True
    return False


_tender_text_cache = {}


def _tender_text(project_id):
    """抽取漏项时回读招标原文；只读 tender 文件，按路径+mtime 缓存。"""
    from app.core.config import settings
    from app.core.database import query
    files = query(
        "SELECT file_path FROM project_file WHERE project_id = %s AND category = 'tender'",
        (project_id,))
    paths = [settings.uploads_dir / f['file_path'] for f in files
             if str(f.get('file_path') or '').lower().endswith(('.docx', '.pdf'))]
    key = (project_id, tuple((str(p), p.stat().st_mtime) for p in paths if p.exists()))
    if key in _tender_text_cache:
        return _tender_text_cache[key]
    from docx import Document
    chunks = []
    for path in paths:
        if not path.exists():
            continue
        try:
            if path.suffix.lower() == '.pdf':
                import pdfplumber
                with pdfplumber.open(str(path)) as pdf:
                    chunks.extend(page.extract_text() or '' for page in pdf.pages)
            else:
                doc = Document(str(path))
                chunks.extend(p.text for p in doc.paragraphs)
                for table in doc.tables:
                    for row in table.rows:
                        chunks.append(' | '.join(c.text for c in row.cells))
        except Exception:  # noqa: BLE001 原文读取失败不影响已抽取的要求
            continue
    text = '\n'.join(chunks)
    _tender_text_cache[key] = text
    return text


def index_required(requirements, project_id=None):
    """只认当前文件的明确编制指令；抽取漏了索引要求时回读招标原文。"""
    for row in requirements:
        if row.get('status') == 'ignored' or row.get('category') != '编制要求':
            continue
        if _text_requires_index(row.get('content')):
            return True
    if project_id is not None:
        return _text_requires_index(_tender_text(project_id))
    return False


def build_entries(requirements, tree):
    nodes = []
    def walk(items):
        for node in items:
            if is_index(node): continue
            nodes.append(node)
            walk(node.get('children') or [])
    walk(tree)
    entries = []
    for req in requirements:
        if req.get('status') == 'ignored':
            continue
        matches = [n for n in nodes if normalize(req['title']) == normalize(n['title'])]
        if not matches:
            matches = [n for n in nodes if related(req['title'], n['title'])]
        if not matches:
            matches = [n for n in nodes if related(req.get('suggested_chapter'), n['title'])]
        refs = []
        seen = set()
        def collect(node, path=()):
            if is_index(node) or node['id'] in seen: return
            seen.add(node['id'])
            refs.append({'id':node['id'], 'title':' / '.join((*path,node['title']))})
            for child in node.get('children') or []:
                collect(child, (*path,node['title']))
        for node in matches:
            collect(node)
        entries.append({'requirement_id': req.get('id'), 'title': req['title'],
                        'content': req.get('content', ''), 'source_file': req.get('source_file', ''),
                        'source_location': req.get('source_location', ''),
                        'node_id': matches[0]['id'] if matches else None,
                        'chapter_title': matches[0]['title'] if matches else '',
                        'chapter_refs': refs,
                        'included': bool(matches)})
    return entries


def project_index(project_id, tree=None):
    from app.services.requirement_service import list_requirements, is_technical_requirement
    from app.services.section_service import get_tree
    requirements = list_requirements(project_id)
    if not index_required(requirements, project_id):
        return []
    # 技术评分索引只列技术评分项；业绩/价格/资格等商务评审项不得进表。
    return build_entries([r for r in requirements
                          if r.get('category') == '评分办法' and is_technical_requirement(r)],
                         get_tree(project_id) if tree is None else tree)
